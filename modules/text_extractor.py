"""
modules/text_extractor.py
Robust text extraction and document statistics for PDF, DOCX, and TXT files.
"""

import re
import math
from typing import Tuple, Dict, Any, List

def extract_text_from_file(uploaded_file) -> Tuple[str, str]:
    """
    Extracts raw text from an uploaded Streamlit file or file-like object.
    Returns (cleaned_text, error_message).
    """
    filename = getattr(uploaded_file, "name", "").lower()
    raw_text = ""
    error = ""

    try:
        if filename.endswith(".pdf"):
            raw_text, error = _extract_pdf(uploaded_file)
        elif filename.endswith(".docx"):
            raw_text, error = _extract_docx(uploaded_file)
        elif filename.endswith(".txt"):
            raw_text, error = _extract_txt(uploaded_file)
        else:
            try:
                raw_text, error = _extract_pdf(uploaded_file)
                if not raw_text:
                    uploaded_file.seek(0)
                    raw_text, error = _extract_docx(uploaded_file)
            except Exception:
                uploaded_file.seek(0)
                raw_text, error = _extract_txt(uploaded_file)
    except Exception as e:
        error = f"Failed to process document: {str(e)}"

    cleaned = clean_text(raw_text)
    return cleaned, error


def _extract_pdf(uploaded_file) -> Tuple[str, str]:
    text_parts = []
    # Try pypdf
    try:
        import pypdf
        reader = pypdf.PdfReader(uploaded_file)
        for page in reader.pages:
            txt = page.extract_text()
            if txt:
                text_parts.append(txt)
        if text_parts:
            return "\n\n".join(text_parts), ""
    except Exception:
        pass

    # Fallback to PyPDF2
    try:
        uploaded_file.seek(0)
        import PyPDF2
        reader = PyPDF2.PdfReader(uploaded_file)
        for page in reader.pages:
            txt = page.extract_text()
            if txt:
                text_parts.append(txt)
        if text_parts:
            return "\n\n".join(text_parts), ""
    except Exception:
        pass

    return "", "Could not extract text from the PDF. It may be scanned or empty."


def _extract_docx(uploaded_file) -> Tuple[str, str]:
    try:
        import docx
        doc = docx.Document(uploaded_file)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        return "\n\n".join(paragraphs), ""
    except Exception as e:
        return "", f"Could not extract text from Word document: {str(e)}"


def _extract_txt(uploaded_file) -> Tuple[str, str]:
    try:
        content = uploaded_file.read()
        for encoding in ["utf-8-sig", "utf-8", "utf-16", "latin-1", "cp1252"]:
            try:
                return content.decode(encoding), ""
            except UnicodeDecodeError:
                continue
        return content.decode("utf-8", errors="replace"), ""
    except Exception as e:
        return "", f"Could not decode text file: {str(e)}"


def clean_text(text: str) -> str:
    """Normalizes excessive whitespace and removes BOM while preserving paragraph structure."""
    if not text:
        return ""
    text = text.lstrip("\ufeff").replace("\ufeff", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\n{3,}", "\n\n", text)
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]
    return "\n".join(lines).strip()


def calculate_document_stats(text: str) -> Dict[str, Any]:
    """Computes comprehensive document statistics."""
    cleaned = clean_text(text)
    if not cleaned:
        return {
            "word_count": 0,
            "char_count": 0,
            "sentence_count": 0,
            "reading_time_min": 0,
            "readability_level": "N/A",
            "paragraph_count": 0,
        }

    words = re.findall(r"\b\w+\b", cleaned)
    word_count = len(words)
    char_count = len(cleaned)

    sentences = re.split(r"(?<=[.!?])\s+", cleaned)
    sentences = [s.strip() for s in sentences if len(s.strip()) > 3]
    sentence_count = max(1, len(sentences))

    reading_time_min = max(1, math.ceil(word_count / 200)) if word_count > 0 else 0
    paragraphs = [p for p in cleaned.split("\n\n") if p.strip()]
    paragraph_count = max(1, len(paragraphs))

    asl = word_count / sentence_count
    fre_score = 206.835 - (1.015 * asl) - (84.6 * 1.45)

    if fre_score >= 80:
        readability_level = "Easy (Elementary)"
    elif fre_score >= 60:
        readability_level = "Standard (High School)"
    elif fre_score >= 40:
        readability_level = "Moderate (College Level)"
    else:
        readability_level = "Advanced (Technical / Academic)"

    return {
        "word_count": word_count,
        "char_count": char_count,
        "sentence_count": sentence_count,
        "reading_time_min": reading_time_min,
        "readability_level": readability_level,
        "paragraph_count": paragraph_count,
    }
